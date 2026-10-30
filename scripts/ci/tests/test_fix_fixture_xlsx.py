# 用途：Sim 配置包 Excel schema 门禁测试。
# 用法：在仓库根目录执行 python3 -m unittest scripts/ci/tests/test_fix_fixture_xlsx.py。
"""验证配置 fixture 不再依赖已退役的 Workflow Job 包装。"""

import importlib.util
import json
import unittest
from pathlib import Path

from openpyxl import Workbook


SCRIPT_PATH = Path(__file__).resolve().parents[2] / "sim" / "fix-fixture-xlsx.py"
SPEC = importlib.util.spec_from_file_location("fix_fixture_xlsx", SCRIPT_PATH)
FIXTURE_PATCHER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(FIXTURE_PATCHER)


class SftpChannelContractTest(unittest.TestCase):
    def test_renamed_keys_match_the_dispatch_consumer(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['channel_type', 'config_json'])
        sheet.append(['SFTP', json.dumps({'host': 'sftp', 'port': 22, 'username': 'ta',
                     'password': 'fixture', 'remotePath': '/inbound'})])
        FIXTURE_PATCHER.patch_channel_config_json(sheet)
        self.assertEqual(json.loads(sheet.cell(2, 2).value), {'sftp_host': 'sftp', 'sftp_port': 22,
                         'sftp_user': 'ta', 'sftp_password': 'fixture', 'sftp_remote_directory': '/inbound'})


class RetireWorkflowJobWrappersTest(unittest.TestCase):
    def setUp(self):
        self.workbook = Workbook()
        jobs = self.workbook.active
        jobs.title = "job_definition"
        jobs.append([
            "tenant_id", "job_code", "job_type", "schedule_type", "schedule_expr",
            "calendar_code", "enabled", "description",
        ])
        jobs.append([
            "ta", "WF_DAILY", "WORKFLOW", "CRON", "0 0 22 * * ?",
            "business-day", "TRUE", "Daily workflow",
        ])
        workflows = self.workbook.create_sheet("workflow_definition")
        workflows.append([
            "tenant_id", "workflow_code", "schedule_enabled", "schedule_type",
            "schedule_expr", "timezone", "trigger_mode", "calendar_code", "description",
        ])
        workflows.append(["ta", "WF_DAILY", None, None, None, None, None, None, None])

    def test_moves_schedule_to_workflow_then_removes_wrapper_job(self):
        retired = FIXTURE_PATCHER.retire_workflow_job_wrappers(self.workbook)

        self.assertEqual(retired, 1)
        self.assertEqual(self.workbook["job_definition"].max_row, 1)
        row = dict(zip(
            [cell.value for cell in self.workbook["workflow_definition"][1]],
            [cell.value for cell in self.workbook["workflow_definition"][2]],
        ))
        self.assertEqual(row["schedule_enabled"], "TRUE")
        self.assertEqual(row["schedule_type"], "CRON")
        self.assertEqual(row["schedule_expr"], "0 0 22 * * ?")
        self.assertEqual(row["timezone"], "Asia/Shanghai")
        self.assertEqual(row["trigger_mode"], "SCHEDULED")
        self.assertEqual(row["calendar_code"], "business-day")

    def test_fails_closed_when_wrapper_has_no_workflow_definition(self):
        self.workbook["workflow_definition"]["B2"] = "OTHER_WORKFLOW"

        with self.assertRaisesRegex(ValueError, "no matching Workflow definition"):
            FIXTURE_PATCHER.retire_workflow_job_wrappers(self.workbook)


if __name__ == "__main__":
    unittest.main()
