# 用途：验证配置包生成器的结构与场景完整性。
# 用法：在仓库根目录执行 python3 -m unittest scripts/ci/tests/test_config_package_fixture_generators.py。
from __future__ import annotations

import importlib.util
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
SUITE_DIR = ROOT / "test-data/tenant-config-packages"
EXPECTED_SHEETS = [
    "resource_queue",
    "business_calendar",
    "batch_window",
    "job_definition",
    "job_monitoring_policy",
    "file_channel_config",
    "file_template_config",
    "workflow_definition",
    "workflow_node",
    "workflow_edge",
]


def load_script(name: str, module_name: str):
    spec = importlib.util.spec_from_file_location(module_name, ROOT / name)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    previous = sys.dont_write_bytecode
    sys.dont_write_bytecode = True
    try:
        spec.loader.exec_module(module)
    finally:
        sys.dont_write_bytecode = previous
    return module


APPEND_SCRIPT = load_script(
    "scripts/local/append-tenant-coverage.py", "append_tenant_coverage"
)
DEFAULT_GENERATOR = load_script(
    "scripts/local/gen-default-tenant-excel.py", "gen_default_tenant_excel"
)


class ConfigPackageFixtureGeneratorsTest(unittest.TestCase):
    def test_new_api_channel_uses_runtime_config_keys(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "tb-tenant-config-package-test.xlsx"
            shutil.copy2(SUITE_DIR / path.name, path)
            workbook = load_workbook(path)
            sheet = workbook["file_channel_config"]
            for row in range(sheet.max_row, 1, -1):
                if sheet.cell(row, 2).value == "tb_api_ingest":
                    sheet.delete_rows(row)
            workbook.save(path)
            workbook.close()
            old_base = APPEND_SCRIPT.BASE
            try:
                APPEND_SCRIPT.BASE = temp_dir + "/"
                APPEND_SCRIPT.update_tb()
            finally:
                APPEND_SCRIPT.BASE = old_base
            workbook = load_workbook(path, read_only=True)
            try:
                row = next(row for row in workbook["file_channel_config"].iter_rows(values_only=True)
                           if row[1] == "tb_api_ingest")
                self.assertEqual({"target_endpoint": row[4]}, json.loads(row[6]))
            finally:
                workbook.close()

    def test_default_generator_emits_current_ten_sheet_contract(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "default-tenant.xlsx"
            DEFAULT_GENERATOR.OUT = str(output)
            DEFAULT_GENERATOR.main()

            workbook = load_workbook(output, read_only=True)

        self.assertEqual(EXPECTED_SHEETS, workbook.sheetnames)

    def test_tenant_coverage_append_is_idempotent_and_does_not_restore_pipeline_sheets(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            suite_dir = Path(temp_dir)
            for tenant in ("ta", "tb", "tc"):
                shutil.copy2(
                    SUITE_DIR / f"{tenant}-tenant-config-package-test.xlsx",
                    suite_dir / f"{tenant}-tenant-config-package-test.xlsx",
                )
            APPEND_SCRIPT.BASE = f"{suite_dir}{Path('/').as_posix()}"

            APPEND_SCRIPT.update_ta()
            APPEND_SCRIPT.update_tb()
            APPEND_SCRIPT.update_tc()
            first_counts = self._tenant_row_counts(suite_dir)
            APPEND_SCRIPT.update_ta()
            APPEND_SCRIPT.update_tb()
            APPEND_SCRIPT.update_tc()
            second_counts = self._tenant_row_counts(suite_dir)

        self.assertEqual(first_counts, second_counts)
        for sheet_names, _ in second_counts.values():
            self.assertEqual(EXPECTED_SHEETS, list(sheet_names))

    @staticmethod
    def _tenant_row_counts(suite_dir: Path) -> dict[str, tuple[tuple[str, ...], tuple[int, int]]]:
        result = {}
        for tenant in ("ta", "tb", "tc"):
            path = suite_dir / f"{tenant}-tenant-config-package-test.xlsx"
            workbook = load_workbook(path, read_only=True, data_only=True)
            result[tenant] = (
                tuple(workbook.sheetnames),
                (
                    workbook["job_definition"].max_row,
                    workbook["file_channel_config"].max_row,
                ),
            )
        return result


if __name__ == "__main__":
    unittest.main()
