# 用途：验证 SIM 导入前的渠道地址归一化及原配置包隔离。
# 用法：python3 -m unittest scripts/ci/tests/test_sim_tenant_fixture.py。
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
SUITE = ROOT / 'test-data/tenant-config-packages'
HELPER = ROOT / 'scripts/sim/prepare-tenant-fixture.py'
IMPORT = ROOT / 'scripts/sim/03-import-tenants.sh'


def channels(path):
    workbook = load_workbook(path, read_only=True)
    try:
        rows = workbook['file_channel_config'].iter_rows(values_only=True)
        header = next(rows)
        return {row[1]: dict(zip(header, row)) for row in rows}
    finally:
        workbook.close()


class SimTenantFixtureTest(unittest.TestCase):
    def prepare(self, source, output, base):
        return subprocess.run([sys.executable, str(HELPER), '--source', str(source),
                               '--output', str(output), '--mockserver-base-url', base],
                              capture_output=True, text=True)

    def test_source_api_channels_use_only_runtime_keys(self):
        for tenant, names in [('tb', ['tb_api_push', 'tb_api_ingest']), ('tc', ['tc_api_risk_push'])]:
            source = SUITE / f'{tenant}-tenant-config-package-test.xlsx'
            for name in names:
                row = channels(source)[name]
                config = json.loads(row['config_json'])
                self.assertFalse({'url', 'method', 'tokenHeader'} & config.keys(), name)
                self.assertEqual(row['target_endpoint'], config['target_endpoint'])

    def test_native_updates_only_target_cells_and_preserves_source(self):
        source = SUITE / 'tb-tenant-config-package-test.xlsx'
        before = source.read_bytes()
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'native.xlsx'
            result = self.prepare(source, output, 'http://localhost:49327')
            self.assertEqual(0, result.returncode, result.stderr)
            prepared = channels(output)
            for name, path in [('tb_api_push', '/tb/callback'), ('tb_api_ingest', '/tb/ingest')]:
                self.assertEqual('http://localhost:49327' + path, prepared[name]['target_endpoint'])
                config = json.loads(prepared[name]['config_json'])
                self.assertEqual(prepared[name]['target_endpoint'], config['target_endpoint'])
                self.assertFalse({'url', 'method', 'tokenHeader'} & config.keys())
            with ZipFile(source) as original, ZipFile(output) as prepared_zip:
                self.assertEqual(original.namelist(), prepared_zip.namelist())
                self.assertEqual(original.comment, prepared_zip.comment)
                changed = [name for name in original.namelist() if original.read(name) != prepared_zip.read(name)]
                self.assertEqual(['xl/worksheets/sheet6.xml'], changed)
                for left, right in zip(original.infolist(), prepared_zip.infolist()):
                    self.assertEqual((left.filename, left.date_time, left.compress_type),
                                     (right.filename, right.date_time, right.compress_type))
            self.assertEqual(0o600, output.stat().st_mode & 0o777)
        self.assertEqual(before, source.read_bytes())

    def test_compose_keeps_template_bytes_and_repeat_hash_is_stable(self):
        for tenant in ('ta', 'tb', 'tc'):
            source = SUITE / f'{tenant}-tenant-config-package-test.xlsx'
            with tempfile.TemporaryDirectory() as directory:
                outputs = [Path(directory) / f'{index}.xlsx' for index in range(2)]
                for output in outputs:
                    result = self.prepare(source, output, 'http://mockserver:1080')
                    self.assertEqual(0, result.returncode, result.stderr)
                self.assertEqual(source.read_bytes(), outputs[0].read_bytes())
                self.assertEqual(outputs[0].read_bytes(), outputs[1].read_bytes())

    def test_native_repeat_hash_and_unknown_channel_are_stable(self):
        source = SUITE / 'tc-tenant-config-package-test.xlsx'
        with tempfile.TemporaryDirectory() as directory:
            first, second = [Path(directory) / f'{name}.xlsx' for name in ('first', 'second')]
            for output in (first, second):
                result = self.prepare(source, output, 'http://127.0.0.1:49328')
                self.assertEqual(0, result.returncode, result.stderr)
            self.assertEqual(first.read_bytes(), second.read_bytes())
            before, after = channels(source), channels(first)
            for name in before.keys() - {'tc_api_risk_push'}:
                self.assertEqual(before[name], after[name], name)

    def test_malformed_known_channel_stops_before_creating_output(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = [Path(directory) / f'{name}.xlsx' for name in ('source', 'output')]
            workbook = load_workbook(SUITE / 'tc-tenant-config-package-test.xlsx')
            workbook['file_channel_config']['G2'] = '{bad-json'
            workbook.save(source)
            workbook.close()
            result = self.prepare(source, output, 'http://localhost:49327')
            self.assertNotEqual(0, result.returncode)
            self.assertFalse(output.exists())

    def run_import(self, recorded=''):
        text = IMPORT.read_text()
        names = ['cleanup_import_artifacts', 'import_one']
        functions = []
        for name in names:
            match = re.search(rf'^{name}\(\) \{{\n.*?^\}}', text, re.M | re.S)
            self.assertIsNotNone(match, name)
            functions.append(match.group())
        with tempfile.TemporaryDirectory() as directory:
            temp = Path(directory)
            work = temp / 'private'
            work.mkdir(mode=0o700)
            capture = temp / 'upload.json'
            mock = temp / 'curl'
            mock.write_text(f'#!{sys.executable}\n' + r'''import hashlib,json,os,sys
from pathlib import Path
from openpyxl import load_workbook
args=sys.argv[1:]; output=Path(args[args.index('-o')+1])
if '-F' in args:
 path=Path(args[args.index('-F')+1].removeprefix('file=@'))
 workbook=load_workbook(path,read_only=True)
 row=next(row for row in workbook['file_channel_config'].iter_rows(values_only=True) if row[1]=='tb_api_push')
 Path(os.environ['CAPTURE']).write_text(json.dumps({'path':str(path),'sha':hashlib.sha256(path.read_bytes()).hexdigest(),'endpoint':row[4],'config':json.loads(row[6])}))
 workbook.close()
 output.write_text('{"code":"SUCCESS","data":{"uploadToken":"fixture-upload"}}')
else: output.write_text('{"code":"SUCCESS"}')
print('200',end='')
''')
            mock.chmod(0o755)
            docker = temp / 'docker'
            docker.write_text(f'#!{sys.executable}\n' + r'''import os,sys
from pathlib import Path
sql=sys.stdin.read(); folder=Path(os.environ['ROOT'])/'scripts/sim/sql'
if sql==(folder/'select-sim-tenant-fixture-import.sql').read_text(): print(os.environ['RECORDED'],end='')
elif sql==(folder/'has-sim-tenant-package-data.sql').read_text(): print('f')
elif sql==(folder/'record-sim-tenant-fixture-import.sql').read_text(): Path(os.environ['RECORD_CAPTURE']).write_text(' '.join(sys.argv[1:]))
else: raise SystemExit('unexpected SQL fixture in transport test')
''')
            docker.chmod(0o755)
            env = dict(os.environ, PATH=str(temp) + os.pathsep + os.environ['PATH'],
                       ROOT=str(ROOT), FIXTURE_DIR=str(SUITE), FIXTURE_WORK_DIR=str(work),
                       PYTHON_BIN=sys.executable, CAPTURE=str(capture), RECORDED=recorded,
                       RECORD_CAPTURE=str(temp / 'record.txt'), PG_CONTAINER='unused',
                       POSTGRES_USER='unused', PLATFORM_DB='unused', CONSOLE_BASE='http://console.invalid')
            script = 'set -euo pipefail\nCURL_AUTH_ARGS=()\nsim_mockserver_base_url() { echo http://localhost:49327; }\n'
            script += '\n'.join(functions) + '\ntrap "cleanup_import_artifacts || exit 1" EXIT\nimport_one tb\n'
            result = subprocess.run(['bash', '-c', script], env=env, capture_output=True, text=True)
            info = json.loads(capture.read_text()) if capture.exists() else None
            record = (temp / 'record.txt').read_text() if (temp / 'record.txt').exists() else ''
            self.assertFalse(work.exists(), '本轮私有副本目录必须清理')
            return result, info, record

    def test_import_upload_and_ledger_use_prepared_copy(self):
        result, info, record = self.run_import()
        self.assertEqual(0, result.returncode, result.stderr)
        self.assertIsNotNone(info)
        self.assertNotEqual(str(SUITE / 'tb-tenant-config-package-test.xlsx'), info['path'])
        self.assertNotEqual(hashlib.sha256((SUITE / 'tb-tenant-config-package-test.xlsx').read_bytes()).hexdigest(), info['sha'])
        self.assertIn('fixture_sha256=' + info['sha'], record)
        self.assertEqual('http://localhost:49327/tb/callback', info['endpoint'])
        self.assertEqual(info['endpoint'], info['config']['target_endpoint'])
        self.assertFalse({'url', 'method', 'tokenHeader'} & info['config'].keys())

    def test_mismatched_ledger_fails_without_upload_and_cleans_copy(self):
        result, info, record = self.run_import('old-fixture-sha')
        self.assertNotEqual(0, result.returncode)
        self.assertIsNone(info)
        self.assertEqual('', record)

    def test_cleanup_on_success_failure_and_signal_is_limited_to_owned_files(self):
        text = IMPORT.read_text()
        cleanup = re.search(r'^cleanup_import_artifacts\(\) \{\n.*?^\}', text, re.M | re.S).group()
        for ending, expected in [('true', 0), ('false', 1), ('kill -TERM $$', 143)]:
            with self.subTest(ending=ending), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                work = root / 'private'
                work.mkdir(mode=0o700)
                for tenant in ('ta', 'tb', 'tc'):
                    (work / f'{tenant}-tenant-config-package-test.xlsx').write_text('owned')
                neighbor = root / 'keep.xlsx'
                neighbor.write_text('unrelated')
                script = 'set -euo pipefail\n' + cleanup
                script += '\ntrap "cleanup_import_artifacts || exit 1" EXIT\ntrap "exit 143" TERM\n' + ending
                result = subprocess.run(['bash', '-c', script], env=dict(os.environ, FIXTURE_WORK_DIR=str(work)), capture_output=True, text=True)
                self.assertEqual(expected, result.returncode, result.stderr)
                self.assertFalse(work.exists())
                self.assertEqual('unrelated', neighbor.read_text())


if __name__ == '__main__':
    unittest.main()
